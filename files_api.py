import hashlib
import os
import shutil
import subprocess
import tempfile
import zipfile
from datetime import datetime, timezone
from pathlib import Path, PurePosixPath
from typing import Annotated
from uuid import uuid4

from fastapi import (
    APIRouter,
    Depends,
    File,
    Form,
    HTTPException,
    Response,
    UploadFile,
    status,
)
from fastapi.responses import FileResponse
from pydantic import BaseModel, Field
from sqlalchemy.orm import Session

from auth import get_current_user
from database import get_db
from ingestion_models import FileIngestionRecord
from models import AuditLog, Document, StoredFile, User
from permissions import require_permission
from schemas import StoredFileListResponse, StoredFileResponse

router = APIRouter()

DBSession = Annotated[Session, Depends(get_db)]
CurrentUser = Annotated[User, Depends(get_current_user)]
Upload = Annotated[UploadFile, File()]

ENVIRONMENT = os.getenv("SAFETY360_ENV", "development").lower()
MAX_FILE_BYTES = int(os.getenv("MAX_FILE_BYTES", str(100 * 1024 * 1024)))
MAX_EXTRACTED_TEXT_CHARS = int(os.getenv("MAX_EXTRACTED_TEXT_CHARS", "2000000"))
MAX_ARCHIVE_ENTRIES = int(os.getenv("MAX_ARCHIVE_ENTRIES", "2000"))
MAX_ARCHIVE_UNCOMPRESSED_BYTES = int(
    os.getenv("MAX_ARCHIVE_UNCOMPRESSED_BYTES", str(256 * 1024 * 1024))
)
MAX_ARCHIVE_COMPRESSION_RATIO = int(os.getenv("MAX_ARCHIVE_COMPRESSION_RATIO", "100"))
STORAGE_PROVIDER = os.getenv("STORAGE_PROVIDER", "local").strip().lower()
STORAGE_ROOT = Path(os.getenv("STORAGE_ROOT", "./storage_data")).resolve()
FILE_SCAN_PROVIDER = os.getenv("FILE_SCAN_PROVIDER", "none").strip().lower()
CLAMSCAN_PATH = os.getenv("CLAMSCAN_PATH", "clamscan").strip()
DOCUMENT_CONVERSION_MODE = os.getenv("DOCUMENT_CONVERSION_MODE", "disabled").strip().lower()
DOCUMENT_CONVERSION_SANDBOXED = (
    os.getenv("DOCUMENT_CONVERSION_SANDBOXED", "false").lower() == "true"
)
LIBREOFFICE_PATH = os.getenv("LIBREOFFICE_PATH", "libreoffice").strip()

SUPPORTED_TYPES: dict[str, tuple[str, str]] = {
    ".pdf": ("pdf", "application/pdf"),
    ".docx": (
        "docx",
        "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
    ),
    ".xlsx": (
        "xlsx",
        "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    ),
    ".pptx": (
        "pptx",
        "application/vnd.openxmlformats-officedocument.presentationml.presentation",
    ),
    ".txt": ("text", "text/plain; charset=utf-8"),
    ".md": ("markdown", "text/markdown; charset=utf-8"),
    ".csv": ("csv", "text/csv; charset=utf-8"),
    ".json": ("json", "application/json"),
}

BLOCKED_EXTENSIONS = {
    ".bat",
    ".cmd",
    ".com",
    ".dll",
    ".docm",
    ".exe",
    ".jar",
    ".js",
    ".msi",
    ".pptm",
    ".ps1",
    ".scr",
    ".sh",
    ".vbs",
    ".xlsm",
}

DANGEROUS_OFFICE_PARTS = (
    "vbaproject.bin",
    "/activex/",
    "/embeddings/",
    "/externallinks/",
    "/oleobjects/",
)


class FileProcessingResponse(BaseModel):
    file_id: int
    detected_format: str
    scan_status: str
    processing_status: str
    parser: str | None = None
    characters: int = 0
    extracted_sha256: str | None = None
    preview: str | None = None
    truncated: bool = False


class ExtractedContentResponse(BaseModel):
    file_id: int
    detected_format: str
    parser: str | None = None
    text: str
    sha256: str


class DocumentFromFileRequest(BaseModel):
    title: str | None = Field(default=None, max_length=250)
    document_type: str = Field(
        default="imported_document",
        min_length=2,
        max_length=80,
    )


class ConvertedFileResponse(BaseModel):
    source_file_id: int
    converted_file: StoredFileResponse


def _require_tenant(current_user: User) -> int:
    if current_user.tenant_id is None:
        raise HTTPException(
            status_code=status.HTTP_409_CONFLICT,
            detail=(
                "Vor der Dateiablage muss ein Mandant angelegt "
                "oder zugeordnet werden."
            ),
        )
    return current_user.tenant_id


def _validate_local_provider() -> None:
    if STORAGE_PROVIDER != "local":
        raise HTTPException(
            status_code=status.HTTP_503_SERVICE_UNAVAILABLE,
            detail=(
                "Der konfigurierte Storage-Provider ist in diesem Build noch nicht "
                "aktiviert. Für Entwicklung und Self-Hosting steht "
                "STORAGE_PROVIDER=local zur Verfügung."
            ),
        )


def _safe_suffix(filename: str | None) -> str:
    suffix = Path(filename or "file").suffix.lower()
    if suffix in BLOCKED_EXTENSIONS:
        raise HTTPException(
            status_code=status.HTTP_415_UNSUPPORTED_MEDIA_TYPE,
            detail="Dieser Dateityp ist aus Sicherheitsgründen nicht zulässig.",
        )
    if suffix not in SUPPORTED_TYPES:
        raise HTTPException(
            status_code=status.HTTP_415_UNSUPPORTED_MEDIA_TYPE,
            detail=(
                "Nicht unterstütztes Dateiformat. Zulässig sind PDF, DOCX, XLSX, "
                "PPTX, TXT, Markdown, CSV und JSON."
            ),
        )
    return suffix


def _normalize_folder(value: str) -> str:
    parts = [
        part
        for part in value.replace("\\", "/").split("/")
        if part and part not in {".", ".."}
    ]
    normalized = "/" + "/".join(parts)
    return normalized[:250] or "/"


def _tenant_file(
    db: Session,
    file_id: int,
    tenant_id: int,
    include_archived: bool = False,
) -> StoredFile:
    query = db.query(StoredFile).filter(
        StoredFile.id == file_id,
        StoredFile.tenant_id == tenant_id,
    )
    if not include_archived:
        query = query.filter(StoredFile.archived_at.is_(None))
    stored_file = query.first()
    if stored_file is None:
        raise HTTPException(status_code=404, detail="Datei wurde nicht gefunden.")
    return stored_file


def _physical_path(storage_key: str) -> Path:
    root = STORAGE_ROOT.resolve()
    candidate = (root / storage_key).resolve()
    if root not in candidate.parents and candidate != root:
        raise HTTPException(
            status_code=status.HTTP_500_INTERNAL_SERVER_ERROR,
            detail="Ungültiger Speicherpfad.",
        )
    return candidate


def _quarantine_path(logical_id: str, suffix: str) -> Path:
    root = STORAGE_ROOT.resolve()
    quarantine = (root / ".quarantine" / f"{logical_id}{suffix}").resolve()
    if root not in quarantine.parents:
        raise HTTPException(status_code=500, detail="Ungültiger Quarantänepfad.")
    quarantine.parent.mkdir(parents=True, exist_ok=True)
    return quarantine


def _validate_office_archive(path: Path, suffix: str) -> None:
    try:
        with zipfile.ZipFile(path) as archive:
            entries = archive.infolist()
            if len(entries) > MAX_ARCHIVE_ENTRIES:
                raise HTTPException(
                    status_code=422,
                    detail="Office-Datei enthält zu viele Archiveinträge.",
                )

            total_uncompressed = 0
            expected_prefix = {
                ".docx": "word/",
                ".xlsx": "xl/",
                ".pptx": "ppt/",
            }[suffix]
            has_expected_part = False

            for entry in entries:
                normalized_name = entry.filename.replace("\\", "/")
                pure = PurePosixPath(normalized_name)
                lowered = f"/{normalized_name.lower()}"
                if pure.is_absolute() or ".." in pure.parts:
                    raise HTTPException(
                        status_code=422,
                        detail="Office-Datei enthält unsichere Archivpfade.",
                    )
                if any(marker in lowered for marker in DANGEROUS_OFFICE_PARTS):
                    raise HTTPException(
                        status_code=422,
                        detail=(
                            "Office-Datei enthält aktive, eingebettete oder externe "
                            "Inhalte und wurde abgelehnt."
                        ),
                    )
                if normalized_name.startswith(expected_prefix):
                    has_expected_part = True
                total_uncompressed += entry.file_size
                if total_uncompressed > MAX_ARCHIVE_UNCOMPRESSED_BYTES:
                    raise HTTPException(
                        status_code=422,
                        detail="Office-Datei überschreitet das Entpack-Limit.",
                    )
                if entry.compress_size > 0:
                    ratio = entry.file_size / entry.compress_size
                    if ratio > MAX_ARCHIVE_COMPRESSION_RATIO:
                        raise HTTPException(
                            status_code=422,
                            detail=(
                                "Office-Datei weist ein unsicheres "
                                "Kompressionsverhältnis auf."
                            ),
                        )

            names = {entry.filename for entry in entries}
            if "[Content_Types].xml" not in names or not has_expected_part:
                raise HTTPException(
                    status_code=422,
                    detail="Office-Dateistruktur passt nicht zur Dateiendung.",
                )
    except zipfile.BadZipFile as exc:
        raise HTTPException(
            status_code=422,
            detail="Office-Datei ist kein gültiges OpenXML-Archiv.",
        ) from exc


def _validate_file_structure(path: Path, suffix: str) -> tuple[str, str]:
    detected_format, media_type = SUPPORTED_TYPES[suffix]
    with path.open("rb") as handle:
        prefix = handle.read(8)

    if suffix == ".pdf":
        if not prefix.startswith(b"%PDF-"):
            raise HTTPException(
                status_code=422,
                detail="PDF-Signatur stimmt nicht mit der Dateiendung überein.",
            )
    elif suffix in {".docx", ".xlsx", ".pptx"}:
        if not prefix.startswith(b"PK\x03\x04"):
            raise HTTPException(
                status_code=422,
                detail="Office-Signatur stimmt nicht mit der Dateiendung überein.",
            )
        _validate_office_archive(path, suffix)
    else:
        try:
            with path.open("r", encoding="utf-8-sig", errors="strict") as handle:
                handle.read(min(MAX_FILE_BYTES, 1024 * 1024))
        except UnicodeDecodeError as exc:
            raise HTTPException(
                status_code=422,
                detail="Textdateien müssen UTF-8-kodiert sein.",
            ) from exc

    return detected_format, media_type


def _scan_file(path: Path) -> str:
    if FILE_SCAN_PROVIDER == "none":
        if ENVIRONMENT == "production":
            raise HTTPException(
                status_code=status.HTTP_503_SERVICE_UNAVAILABLE,
                detail=(
                    "Dateiupload ist in Produktion ohne aktivierten "
                    "Malware-Scanner gesperrt."
                ),
            )
        return "not_configured"

    if FILE_SCAN_PROVIDER != "clamav":
        raise HTTPException(
            status_code=status.HTTP_503_SERVICE_UNAVAILABLE,
            detail="Unbekannter Malware-Scan-Provider.",
        )

    try:
        completed = subprocess.run(
            [CLAMSCAN_PATH, "--no-summary", str(path)],
            capture_output=True,
            text=True,
            timeout=60,
            check=False,
        )
    except (OSError, subprocess.TimeoutExpired) as exc:
        raise HTTPException(
            status_code=status.HTTP_503_SERVICE_UNAVAILABLE,
            detail="Malware-Scanner ist nicht verfügbar.",
        ) from exc

    if completed.returncode == 0:
        return "clean"
    if completed.returncode == 1:
        raise HTTPException(
            status_code=422,
            detail="Datei wurde vom Malware-Scanner abgelehnt.",
        )
    raise HTTPException(
        status_code=status.HTTP_503_SERVICE_UNAVAILABLE,
        detail="Malware-Scan konnte nicht zuverlässig abgeschlossen werden.",
    )


def _register_ingestion(
    db: Session,
    stored_file: StoredFile,
    detected_format: str,
    scan_status: str,
    created_by_id: int,
) -> FileIngestionRecord:
    record = FileIngestionRecord(
        file_id=stored_file.id,
        tenant_id=stored_file.tenant_id,
        detected_format=detected_format,
        scan_status=scan_status,
        processing_status="stored",
        source_sha256=stored_file.sha256,
        created_by_id=created_by_id,
    )
    db.add(record)
    return record


def _extract_text(path: Path, detected_format: str) -> tuple[str, str, bool]:
    chunks: list[str] = []
    current = 0
    truncated = False

    def add(value: object) -> bool:
        nonlocal current, truncated
        text = str(value).strip()
        if not text:
            return True
        remaining = MAX_EXTRACTED_TEXT_CHARS - current
        if remaining <= 0:
            truncated = True
            return False
        piece = text[:remaining]
        chunks.append(piece)
        current += len(piece) + 1
        if len(text) > remaining:
            truncated = True
            return False
        return True

    if detected_format == "pdf":
        import pdfplumber

        with pdfplumber.open(path) as pdf:
            for page in pdf.pages:
                if not add(page.extract_text() or ""):
                    break
        parser = "pdfplumber"
    elif detected_format == "docx":
        from docx import Document as WordDocument

        document = WordDocument(str(path))
        for paragraph in document.paragraphs:
            if not add(paragraph.text):
                break
        if not truncated:
            for table in document.tables:
                for row in table.rows:
                    if not add(" | ".join(cell.text for cell in row.cells)):
                        break
                if truncated:
                    break
        parser = "python-docx"
    elif detected_format == "xlsx":
        from openpyxl import load_workbook

        workbook = load_workbook(filename=path, read_only=True, data_only=True)
        try:
            for worksheet in workbook.worksheets:
                if not add(f"[{worksheet.title}]"):
                    break
                for row in worksheet.iter_rows(values_only=True):
                    values = ["" if value is None else str(value) for value in row]
                    if not add("\t".join(values)):
                        break
                if truncated:
                    break
        finally:
            workbook.close()
        parser = "openpyxl"
    elif detected_format == "pptx":
        from pptx import Presentation

        presentation = Presentation(str(path))
        for index, slide in enumerate(presentation.slides, start=1):
            if not add(f"[Slide {index}]"):
                break
            for shape in slide.shapes:
                if hasattr(shape, "text") and not add(shape.text):
                    break
            if truncated:
                break
        parser = "python-pptx"
    elif detected_format in {"text", "markdown", "csv", "json"}:
        with path.open("r", encoding="utf-8-sig", errors="strict") as handle:
            while True:
                block = handle.read(64 * 1024)
                if not block:
                    break
                if not add(block):
                    break
        parser = "utf8-text"
    else:
        raise HTTPException(
            status_code=status.HTTP_415_UNSUPPORTED_MEDIA_TYPE,
            detail="Für dieses Format ist keine Textextraktion verfügbar.",
        )

    text = "\n".join(chunks).strip()
    return text, parser, truncated


def _process_file_record(
    db: Session,
    stored_file: StoredFile,
    record: FileIngestionRecord,
    current_user: User,
) -> FileProcessingResponse:
    path = _physical_path(stored_file.storage_key)
    if not path.is_file():
        raise HTTPException(
            status_code=status.HTTP_503_SERVICE_UNAVAILABLE,
            detail="Dateiinhalt ist im Speicher nicht verfügbar.",
        )

    record.processing_status = "processing"
    record.error_message = None
    db.commit()

    try:
        text, parser, truncated = _extract_text(path, record.detected_format)
    except HTTPException:
        record.processing_status = "failed"
        record.error_message = "Dokumentverarbeitung wurde abgelehnt."
        db.commit()
        raise
    except Exception as exc:
        record.processing_status = "failed"
        record.error_message = "Dokument konnte nicht sicher verarbeitet werden."
        db.commit()
        raise HTTPException(
            status_code=status.HTTP_422_UNPROCESSABLE_ENTITY,
            detail="Dokument konnte nicht sicher verarbeitet werden.",
        ) from exc

    record.extracted_text = text
    record.extracted_sha256 = hashlib.sha256(text.encode("utf-8")).hexdigest()
    record.parser = parser
    record.processing_status = "processed"
    db.add(
        AuditLog(
            event=f"file_processed:{stored_file.id}:{record.extracted_sha256}",
            user_id=current_user.id,
            tenant_id=stored_file.tenant_id,
        )
    )
    db.commit()
    db.refresh(record)
    return FileProcessingResponse(
        file_id=stored_file.id,
        detected_format=record.detected_format,
        scan_status=record.scan_status,
        processing_status=record.processing_status,
        parser=record.parser,
        characters=len(text),
        extracted_sha256=record.extracted_sha256,
        preview=text[:2000] or None,
        truncated=truncated,
    )


def _persist_generated_file(
    db: Session,
    source_path: Path,
    original_name: str,
    tenant_id: int,
    created_by_id: int,
    category: str,
    folder: str,
) -> StoredFile:
    suffix = _safe_suffix(original_name)
    detected_format, media_type = _validate_file_structure(source_path, suffix)
    scan_status = _scan_file(source_path)
    logical_id = str(uuid4())
    storage_key = f"tenant-{tenant_id}/{logical_id}{suffix}"
    target = _physical_path(storage_key)
    target.parent.mkdir(parents=True, exist_ok=True)
    shutil.copy2(source_path, target)

    hasher = hashlib.sha256()
    size = 0
    with target.open("rb") as handle:
        while True:
            chunk = handle.read(1024 * 1024)
            if not chunk:
                break
            size += len(chunk)
            hasher.update(chunk)

    stored_file = StoredFile(
        logical_id=logical_id,
        tenant_id=tenant_id,
        original_name=original_name[:255],
        storage_key=storage_key,
        media_type=media_type,
        size_bytes=size,
        sha256=hasher.hexdigest(),
        category=category[:80],
        folder=_normalize_folder(folder),
        created_by_id=created_by_id,
    )
    db.add(stored_file)
    db.flush()
    _register_ingestion(
        db,
        stored_file,
        detected_format,
        scan_status,
        created_by_id,
    )
    return stored_file


@router.post(
    "",
    response_model=StoredFileResponse,
    status_code=status.HTTP_201_CREATED,
)
async def upload_file(
    file: Upload,
    current_user: CurrentUser,
    db: DBSession,
    category: Annotated[str, Form(max_length=80)] = "general",
    folder: Annotated[str, Form(max_length=250)] = "/",
) -> StoredFile:
    require_permission(current_user, "files.upload")
    tenant_id = _require_tenant(current_user)
    _validate_local_provider()

    original_name = (file.filename or "file").strip()[:255]
    suffix = _safe_suffix(original_name)
    logical_id = str(uuid4())
    quarantine = _quarantine_path(logical_id, suffix)
    hasher = hashlib.sha256()
    size = 0

    try:
        with quarantine.open("wb") as handle:
            while True:
                chunk = await file.read(1024 * 1024)
                if not chunk:
                    break
                size += len(chunk)
                if size > MAX_FILE_BYTES:
                    raise HTTPException(
                        status_code=status.HTTP_413_REQUEST_ENTITY_TOO_LARGE,
                        detail=(
                            "Die hochgeladene Datei überschreitet das "
                            "konfigurierte Größenlimit."
                        ),
                    )
                hasher.update(chunk)
                handle.write(chunk)

        detected_format, media_type = _validate_file_structure(quarantine, suffix)
        scan_status = _scan_file(quarantine)

        storage_key = f"tenant-{tenant_id}/{logical_id}{suffix}"
        target = _physical_path(storage_key)
        target.parent.mkdir(parents=True, exist_ok=True)
        os.replace(quarantine, target)
    except Exception:
        quarantine.unlink(missing_ok=True)
        raise
    finally:
        await file.close()

    stored_file = StoredFile(
        logical_id=logical_id,
        tenant_id=tenant_id,
        original_name=original_name,
        storage_key=storage_key,
        media_type=media_type,
        size_bytes=size,
        sha256=hasher.hexdigest(),
        category=(category.strip().lower() or "general")[:80],
        folder=_normalize_folder(folder),
        created_by_id=current_user.id,
    )
    db.add(stored_file)
    db.flush()
    _register_ingestion(
        db,
        stored_file,
        detected_format,
        scan_status,
        current_user.id,
    )
    db.add(
        AuditLog(
            event=(
                f"file_uploaded:{stored_file.id}:{stored_file.sha256}:"
                f"{scan_status}"
            ),
            user_id=current_user.id,
            tenant_id=tenant_id,
        )
    )
    db.commit()
    db.refresh(stored_file)
    return stored_file


@router.get("", response_model=StoredFileListResponse)
def list_files(
    current_user: CurrentUser,
    db: DBSession,
    include_archived: bool = False,
) -> StoredFileListResponse:
    require_permission(current_user, "files.read")
    tenant_id = _require_tenant(current_user)
    query = db.query(StoredFile).filter(StoredFile.tenant_id == tenant_id)
    if not include_archived:
        query = query.filter(StoredFile.archived_at.is_(None))
    files = query.order_by(StoredFile.created_at.desc()).all()
    return StoredFileListResponse(files=files)


@router.get("/{file_id}/download")
def download_file(
    file_id: int,
    current_user: CurrentUser,
    db: DBSession,
):
    require_permission(current_user, "files.read")
    tenant_id = _require_tenant(current_user)
    _validate_local_provider()
    stored_file = _tenant_file(db, file_id, tenant_id)
    path = _physical_path(stored_file.storage_key)
    if not path.is_file():
        raise HTTPException(
            status_code=status.HTTP_503_SERVICE_UNAVAILABLE,
            detail="Dateiinhalt ist im Speicher nicht verfügbar.",
        )

    db.add(
        AuditLog(
            event=f"file_downloaded:{stored_file.id}",
            user_id=current_user.id,
            tenant_id=tenant_id,
        )
    )
    db.commit()

    return FileResponse(
        path,
        media_type=stored_file.media_type or "application/octet-stream",
        filename=stored_file.original_name,
        content_disposition_type="attachment",
        headers={
            "X-Content-Type-Options": "nosniff",
            "Cache-Control": "private, no-store",
        },
    )


@router.post("/{file_id}/process", response_model=FileProcessingResponse)
def process_file(
    file_id: int,
    current_user: CurrentUser,
    db: DBSession,
) -> FileProcessingResponse:
    require_permission(current_user, "files.process")
    tenant_id = _require_tenant(current_user)
    stored_file = _tenant_file(db, file_id, tenant_id)
    record = (
        db.query(FileIngestionRecord)
        .filter(
            FileIngestionRecord.file_id == stored_file.id,
            FileIngestionRecord.tenant_id == tenant_id,
        )
        .first()
    )
    if record is None:
        raise HTTPException(
            status_code=status.HTTP_409_CONFLICT,
            detail="Für die Datei fehlt ein sicherer Ingestion-Datensatz.",
        )
    return _process_file_record(db, stored_file, record, current_user)


@router.get("/{file_id}/content", response_model=ExtractedContentResponse)
def extracted_content(
    file_id: int,
    current_user: CurrentUser,
    db: DBSession,
) -> ExtractedContentResponse:
    require_permission(current_user, "files.read")
    tenant_id = _require_tenant(current_user)
    stored_file = _tenant_file(db, file_id, tenant_id)
    record = (
        db.query(FileIngestionRecord)
        .filter(
            FileIngestionRecord.file_id == stored_file.id,
            FileIngestionRecord.tenant_id == tenant_id,
        )
        .first()
    )
    if (
        record is None
        or record.processing_status != "processed"
        or record.extracted_text is None
    ):
        raise HTTPException(
            status_code=status.HTTP_409_CONFLICT,
            detail="Datei wurde noch nicht erfolgreich verarbeitet.",
        )
    return ExtractedContentResponse(
        file_id=stored_file.id,
        detected_format=record.detected_format,
        parser=record.parser,
        text=record.extracted_text,
        sha256=(
            record.extracted_sha256
            or hashlib.sha256(record.extracted_text.encode("utf-8")).hexdigest()
        ),
    )


@router.post(
    "/{file_id}/to-document",
    status_code=status.HTTP_201_CREATED,
)
def create_controlled_document_from_file(
    file_id: int,
    data: DocumentFromFileRequest,
    current_user: CurrentUser,
    db: DBSession,
):
    require_permission(current_user, "files.process")
    require_permission(current_user, "documents.create")
    tenant_id = _require_tenant(current_user)
    stored_file = _tenant_file(db, file_id, tenant_id)
    record = (
        db.query(FileIngestionRecord)
        .filter(
            FileIngestionRecord.file_id == file_id,
            FileIngestionRecord.tenant_id == tenant_id,
        )
        .first()
    )
    if (
        record is None
        or record.processing_status != "processed"
        or record.extracted_text is None
    ):
        raise HTTPException(
            status_code=status.HTTP_409_CONFLICT,
            detail="Datei muss vor der Übernahme zuerst verarbeitet werden.",
        )

    title = (data.title or Path(stored_file.original_name).stem).strip()[:250]
    document = Document(
        logical_id=str(uuid4()),
        title=title,
        document_type=data.document_type.strip().lower(),
        status="draft",
        version=1,
        content_summary=record.extracted_text[:20000],
        tenant_id=tenant_id,
        created_by_id=current_user.id,
    )
    db.add(document)
    db.flush()
    db.add(
        AuditLog(
            event=f"file_promoted_to_document:{stored_file.id}:{document.id}",
            user_id=current_user.id,
            tenant_id=tenant_id,
        )
    )
    db.commit()
    db.refresh(document)
    return {
        "document_id": document.id,
        "logical_id": document.logical_id,
        "status": document.status,
        "version": document.version,
        "source_file_id": stored_file.id,
        "source_sha256": stored_file.sha256,
        "note": (
            "Originaldatei bleibt unverändert erhalten; Bearbeitung erfolgt "
            "revisionssicher im Dokumentenworkflow."
        ),
    }


@router.post(
    "/{file_id}/convert/pdf",
    response_model=ConvertedFileResponse,
)
def convert_file_to_pdf(
    file_id: int,
    current_user: CurrentUser,
    db: DBSession,
) -> ConvertedFileResponse:
    require_permission(current_user, "files.process")
    tenant_id = _require_tenant(current_user)
    stored_file = _tenant_file(db, file_id, tenant_id)
    source_path = _physical_path(stored_file.storage_key)
    suffix = Path(stored_file.original_name).suffix.lower()

    if suffix == ".pdf":
        raise HTTPException(
            status_code=status.HTTP_409_CONFLICT,
            detail="Datei liegt bereits als PDF vor.",
        )
    if suffix not in {".docx", ".xlsx", ".pptx"}:
        raise HTTPException(
            status_code=status.HTTP_415_UNSUPPORTED_MEDIA_TYPE,
            detail="PDF-Konvertierung ist für dieses Format nicht verfügbar.",
        )
    if DOCUMENT_CONVERSION_MODE != "libreoffice":
        raise HTTPException(
            status_code=status.HTTP_503_SERVICE_UNAVAILABLE,
            detail="Dokumentkonvertierung ist nicht aktiviert.",
        )
    if ENVIRONMENT == "production" and not DOCUMENT_CONVERSION_SANDBOXED:
        raise HTTPException(
            status_code=status.HTTP_503_SERVICE_UNAVAILABLE,
            detail=(
                "Produktive Office-Konvertierung ist ohne isolierten "
                "Sandbox-Worker gesperrt."
            ),
        )

    with tempfile.TemporaryDirectory(prefix="safety360-convert-") as temp_dir:
        workdir = Path(temp_dir)
        safe_input = workdir / f"source{suffix}"
        shutil.copy2(source_path, safe_input)
        try:
            completed = subprocess.run(
                [
                    LIBREOFFICE_PATH,
                    "--headless",
                    "--convert-to",
                    "pdf",
                    "--outdir",
                    str(workdir),
                    str(safe_input),
                ],
                capture_output=True,
                text=True,
                timeout=90,
                check=False,
                cwd=workdir,
                env={**os.environ, "HOME": str(workdir)},
            )
        except (OSError, subprocess.TimeoutExpired) as exc:
            raise HTTPException(
                status_code=status.HTTP_503_SERVICE_UNAVAILABLE,
                detail="Dokumentkonverter ist nicht verfügbar.",
            ) from exc

        output = workdir / "source.pdf"
        if completed.returncode != 0 or not output.is_file():
            raise HTTPException(
                status_code=status.HTTP_422_UNPROCESSABLE_ENTITY,
                detail="Dokument konnte nicht in PDF umgewandelt werden.",
            )

        converted_name = f"{Path(stored_file.original_name).stem}.pdf"
        converted = _persist_generated_file(
            db=db,
            source_path=output,
            original_name=converted_name,
            tenant_id=tenant_id,
            created_by_id=current_user.id,
            category="converted",
            folder=stored_file.folder,
        )
        db.add(
            AuditLog(
                event=f"file_converted_to_pdf:{stored_file.id}:{converted.id}",
                user_id=current_user.id,
                tenant_id=tenant_id,
            )
        )
        db.commit()
        db.refresh(converted)

    return ConvertedFileResponse(
        source_file_id=stored_file.id,
        converted_file=StoredFileResponse.model_validate(converted),
    )


@router.post("/{file_id}/archive", response_model=StoredFileResponse)
def archive_file(
    file_id: int,
    current_user: CurrentUser,
    db: DBSession,
) -> StoredFile:
    require_permission(current_user, "files.archive")
    tenant_id = _require_tenant(current_user)
    stored_file = _tenant_file(db, file_id, tenant_id)
    stored_file.archived_at = datetime.now(timezone.utc)
    db.add(
        AuditLog(
            event=f"file_archived:{stored_file.id}",
            user_id=current_user.id,
            tenant_id=tenant_id,
        )
    )
    db.commit()
    db.refresh(stored_file)
    return stored_file


@router.post("/{file_id}/restore", response_model=StoredFileResponse)
def restore_file(
    file_id: int,
    current_user: CurrentUser,
    db: DBSession,
) -> StoredFile:
    require_permission(current_user, "files.archive")
    tenant_id = _require_tenant(current_user)
    stored_file = _tenant_file(
        db,
        file_id,
        tenant_id,
        include_archived=True,
    )
    stored_file.archived_at = None
    db.add(
        AuditLog(
            event=f"file_restored:{stored_file.id}",
            user_id=current_user.id,
            tenant_id=tenant_id,
        )
    )
    db.commit()
    db.refresh(stored_file)
    return stored_file


@router.delete("/{file_id}", status_code=status.HTTP_204_NO_CONTENT)
def delete_file_permanently(
    file_id: int,
    current_user: CurrentUser,
    db: DBSession,
) -> Response:
    require_permission(current_user, "files.delete")
    tenant_id = _require_tenant(current_user)
    stored_file = _tenant_file(
        db,
        file_id,
        tenant_id,
        include_archived=True,
    )
    if stored_file.archived_at is None:
        raise HTTPException(
            status_code=status.HTTP_409_CONFLICT,
            detail="Datei muss vor einer endgültigen Löschung archiviert werden.",
        )

    path = _physical_path(stored_file.storage_key)
    file_hash = stored_file.sha256
    db.add(
        AuditLog(
            event=f"file_permanently_deleted:{stored_file.id}:{file_hash}",
            user_id=current_user.id,
            tenant_id=tenant_id,
        )
    )
    ingestion = (
        db.query(FileIngestionRecord)
        .filter(
            FileIngestionRecord.file_id == stored_file.id,
            FileIngestionRecord.tenant_id == tenant_id,
        )
        .first()
    )
    if ingestion is not None:
        db.delete(ingestion)
    db.delete(stored_file)
    db.commit()
    path.unlink(missing_ok=True)
    return Response(status_code=status.HTTP_204_NO_CONTENT)
